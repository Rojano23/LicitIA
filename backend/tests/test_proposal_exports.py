"""M9C-01–33 synthetic proposal exports; real local writers, no blind/provider runs."""
from copy import deepcopy
from datetime import datetime,timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile,ZIP_DEFLATED
import json
import pytest
from sqlalchemy import select,event
from app.models import Requirement,RequirementEvidenceExpectation,ProposalComposition
from app.proposal_export_contract import ExportRequest
from app.proposal_exports import stable_content,verify_package,safe_name,_base
from app.proposal_export_renderers import RENDERERS
from tests.test_workspace_flow import scene
from tests.test_bid_readiness import complete
from tests.test_proposal_readiness import proposal_connection,create as create_snapshot
from tests.test_proposal_versions import ready,create
from tests.test_proposal_composition import finish,init,change,route,synthetic_snapshot
from tests.test_commercial_matrix import matrix,bulk
from tests.test_readiness_work import read


def freeze(c,t,v,comp,**kwargs):
    response=c.post(route(t,v)+'/freeze',json={'expected_edit_revision':v['edit_revision'],
        'expected_composition_revision':comp['revision'],'acknowledge_warnings':True,'acknowledge_stale':True,**kwargs})
    assert response.status_code==200,response.text
    return response.json()


@pytest.fixture
def proposal(ready):
    c,t,db,m,snapshot=ready;v,_=create(c,t,snapshot);comp=finish(c,t,v);v=freeze(c,t,v,comp)
    return c,t,db,m,snapshot,v,comp


def generate(c,t,v,**changes):
    payload={'request_id':str(uuid4()),**changes}
    response=c.post(route(t,v)+'/exports',json=payload);assert response.status_code==200,response.text
    return response.json(),payload


def artifact(c,t,v,receipt,kind):
    response=c.get(route(t,v)+f"/exports/{receipt['request_id']}/files/{kind}")
    assert response.status_code==200,response.text
    return response.content


def visible_docx(data):
    from docx import Document
    document=Document(BytesIO(data))
    return '\n'.join([p.text for p in document.paragraphs]+[p.text for section in document.sections for p in section.header.paragraphs]+[cell.text for table in document.tables for row in table.rows for cell in row.cells])


def visible_xlsx(data):
    from openpyxl import load_workbook
    book=load_workbook(BytesIO(data))
    return book,'\n'.join(str(cell.value) for sheet in book if sheet.sheet_state=='visible' for row in sheet for cell in row if cell.value is not None)


def visible_pdf(data):
    import pymupdf
    with pymupdf.open(stream=data,filetype='pdf') as pdf:return '\n'.join(page.get_text() for page in pdf)


def test_m9c01_02_03_04_05_06_09_10_11_18_30_31_exact_outputs(proposal):
    c,t,db,m,snapshot,v,comp=proposal
    before=matrix(c,m);snapshot_before=read(c,f"/tenders/{t}/proposal-readiness/snapshots/{v['snapshot_id']}")
    receipt,_=generate(c,t,v)
    data={kind:artifact(c,t,v,receipt,kind) for kind in ['DOCX','XLSX','PDF','ZIP']}
    manifest=verify_package(data['ZIP']);content=manifest['content']
    assert manifest['version_id']==v['id'] and manifest['snapshot_sha256']==snapshot['payload_sha256']
    assert manifest['composition_sha256']==comp['sha256'] and receipt['proposal_state']=='FROZEN'
    assert {a['type'] for a in receipt['artifacts']}=={'DOCX','XLSX','PDF','ZIP'}
    word=visible_docx(data['DOCX']);book,excel=visible_xlsx(data['XLSX']);pdf=visible_pdf(data['PDF'])
    assert 'FROZEN' in word and 'FROZEN' in excel and 'FROZEN' in pdf
    for row in snapshot['payload']['projection']['rows']:
        assert row['description'] in word and row['description'] in excel and row['description'] in pdf
        assert row['offered_unit_price'] in word and row['offered_total'] in word
        assert row['offered_total'] in pdf
    economic=book['Propuesta Economica']
    from decimal import Decimal
    expected={(r['description'],Decimal(r['quantity']),Decimal(r['offered_unit_price']),Decimal(r['offered_total']),r['currency']) for r in snapshot['payload']['projection']['rows']}
    actual={(r[1].value,Decimal(str(r[2].value)),Decimal(str(r[4].value)),Decimal(str(r[5].value)),r[6].value) for r in economic.iter_rows(min_row=4)}
    assert actual==expected and book['_Metadata'].sheet_state=='hidden'
    assert v['id'] not in word and v['snapshot_id'] not in word
    assert c.post(route(t,v)+f"/exports/{receipt['request_id']}/verify").json()['verification_status']=='VERIFIED'
    assert matrix(c,m)==before and read(c,f"/tenders/{t}/proposal-readiness/snapshots/{v['snapshot_id']}")==snapshot_before
    for kind in ['DOCX','XLSX','PDF']:
        identity=next(a for a in receipt['artifacts'] if a['type']==kind)
        assert identity['sha256'] and identity['created_at']
    from docx import Document
    import pymupdf
    assert receipt['semantic_sha256'] in Document(BytesIO(data['DOCX'])).core_properties.comments
    assert json.loads(book.properties.description)['semantic_sha256']==receipt['semantic_sha256']
    with pymupdf.open(stream=data['PDF'],filetype='pdf') as document:assert json.loads(document.metadata['subject'])['semantic_sha256']==receipt['semantic_sha256']
    imported=c.post(f'/tenders/{t}/commercial-exchange/stage',files={'file':('proposal.xlsx',data['XLSX'],'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert imported.status_code==422,imported.text
    assert content['composition_contract']=='proposal-composition-v1' and manifest['contract_version']=='proposal-export-v1'


def test_m9c07_08_19_32_33_frozen_text_survives_current_source_change(proposal,monkeypatch):
    c,t,db,m,snapshot,v,comp=proposal
    requirement=db.scalar(select(Requirement).where(Requirement.tender_id==t));evidence=db.scalar(select(RequirementEvidenceExpectation).where(RequirementEvidenceExpectation.requirement_id==requirement.id))
    old_requirement=requirement.canonical_text;old_evidence=evidence.evidence_description
    requirement.canonical_text='NEW_REQUIREMENT_NOT_FOR_OLD_EXPORT';evidence.evidence_description='NEW_EVIDENCE_NOT_FOR_OLD_EXPORT';db.flush()
    assert read(c,route(t,v)+'/source-state')['state']=='SOURCE_SNAPSHOT_STALE'
    def denied(*a,**k):raise AssertionError('Export must not query current Requirements or network')
    monkeypatch.setattr('app.requirement_matrix.get_tender_requirement_matrix',denied)
    receipt,_=generate(c,t,v)
    output=[visible_docx(artifact(c,t,v,receipt,'DOCX')),visible_xlsx(artifact(c,t,v,receipt,'XLSX'))[1],visible_pdf(artifact(c,t,v,receipt,'PDF'))]
    for text in output:
        assert ' '.join(old_requirement.split()) in ' '.join(text.split()) and ' '.join(old_evidence.split()) in ' '.join(text.split())
        assert 'NEW_REQUIREMENT_NOT_FOR_OLD_EXPORT' not in text and 'NEW_EVIDENCE_NOT_FOR_OLD_EXPORT' not in text
    assert receipt['snapshot_id']==snapshot['payload']['snapshot_id']


def test_m9c12_mixed_currencies_no_aggregation(ready):
    c,t,db,m,snapshot=ready
    bulk(c,m,[matrix(c,m)['rows'][0]],changes={'currency':'USD'},apply=True,confirm_currency_change=True)
    newer,_=create_snapshot(c,t);v,_=create(c,t,newer);comp=finish(c,t,v);v=freeze(c,t,v,comp)
    receipt,_=generate(c,t,v);manifest=verify_package(artifact(c,t,v,receipt,'ZIP'))
    economic=next(s for s in manifest['content']['sections'] if s['id']=='economic')
    assert {row['group'] for row in economic['rows']}=={'MXN','USD'}
    assert all(set(row['values'])=={'item','description','quantity','unit','unit_price','total','currency'} for row in economic['rows'])
    assert 'Moneda: USD' in visible_docx(artifact(c,t,v,receipt,'DOCX'))
    assert 'Moneda: MXN' in visible_pdf(artifact(c,t,v,receipt,'PDF'))


def test_m9c13_14_15_16_internal_cost_notes_absent(ready):
    c,t,db,m,snapshot=ready
    bulk(c,m,matrix(c,m)['rows'],changes={'unit_cost':'987654.32','notes':'PRIVATE_M7_PRICING_STRATEGY'},apply=True)
    newer,_=create_snapshot(c,t);v,_=create(c,t,newer);comp=finish(c,t,v)
    econ=next(r for r in comp['rows'] if r['section']=='ECONOMIC')
    comp=change(c,t,v,comp,econ['key'],{'note':'PRIVATE_M9B_ECONOMIC_NOTE'});v=freeze(c,t,v,comp)
    receipt,_=generate(c,t,v)
    files={kind:artifact(c,t,v,receipt,kind) for kind in ['DOCX','XLSX','PDF','ZIP']}
    exposed=visible_docx(files['DOCX'])+visible_xlsx(files['XLSX'])[1]+visible_pdf(files['PDF'])+json.dumps(verify_package(files['ZIP']))
    for kind in ['DOCX','XLSX']:
        with ZipFile(BytesIO(files[kind])) as z:exposed+=''.join(z.read(name).decode(errors='ignore') for name in z.namelist() if name.endswith('.xml'))
    for secret in ['987654.32','PRIVATE_M7_PRICING_STRATEGY','PRIVATE_M9B_ECONOMIC_NOTE','unit_cost','markup_percent','target_margin_percent']:
        assert secret not in exposed


def test_m9c17_draft_explicit_visible_and_history_remains_draft(ready):
    c,t,db,m,snapshot=ready;v,_=create(c,t,snapshot);comp=init(c,t,v)
    assert c.post(route(t,v)+'/exports',json={'request_id':str(uuid4())}).status_code==409
    receipt,payload=generate(c,t,v,allow_draft=True,expected_composition_revision=comp['revision'],expected_edit_revision=v['edit_revision'])
    assert receipt['proposal_state']=='DRAFT'
    assert 'DRAFT' in visible_docx(artifact(c,t,v,receipt,'DOCX'))
    book,text=visible_xlsx(artifact(c,t,v,receipt,'XLSX'));assert all('DRAFT' in str(s['A1'].value) for s in book if s.sheet_state=='visible')
    assert 'DRAFT' in visible_pdf(artifact(c,t,v,receipt,'PDF'))
    comp=finish(c,t,v);frozen=freeze(c,t,v,comp)
    assert c.get(route(t,v)+f"/exports/{receipt['request_id']}").json()['proposal_state']=='DRAFT'
    assert 'DRAFT' in visible_docx(artifact(c,t,frozen,receipt,'DOCX'))


def test_m9c20_21_26_paths_ownership_and_safe_names(proposal,tmp_path):
    c,t,db,m,snapshot,v,comp=proposal;receipt,_=generate(c,t,v,formats=['DOCX'],include_package=False)
    other=c.post('/tenders',json={'title':'Other'}).json()['id']
    for suffix in ['',f"/{receipt['request_id']}",f"/{receipt['request_id']}/files/DOCX"]:
        assert c.get(route(other,v)+'/exports'+suffix).status_code==404
    assert c.post(route(t,v)+'/exports',json={'request_id':'../../outside'}).status_code==422
    assert c.post(route(t,v)+'/exports',json={'request_id':str(uuid4()),'output_path':'../../outside'}).status_code==422
    assert c.get(route(t,v)+f"/exports/{receipt['request_id']}/files/../../outside").status_code!=200
    name=safe_name('../../bad\r\nname: / 🙂',1,'zip')
    assert '/' not in name and '\\' not in name and '\r' not in name and '..' not in name
    base=_base(t,v['id']);outside=tmp_path/'outside';outside.mkdir();link=base/str(uuid4());link.symlink_to(outside,target_is_directory=True)
    assert c.get(route(t,v)+f'/exports/{link.name}').status_code==409


def test_m9c22_23_24_25_formula_macros_templates_and_remote_fetch(proposal,monkeypatch):
    c,t,db,m,snapshot,v,comp=proposal
    content=stable_content(db,t,v['id'],ExportRequest(request_id=uuid4()))
    dangerous=['=HYPERLINK("https://invalid.test", "x")','+SUM(A1:A2)','-cmd|x','@SUM(A1:A2)','{{7*7}}','<img src="https://invalid.test/a.png"/>']
    section=next(s for s in content.sections if s.id=='requirements')
    keys=list(section.rows[0].values)
    for index,text in enumerate(dangerous):section.rows[0].values[keys[index]]=text
    def no_network(*a,**k):raise AssertionError('No external network')
    monkeypatch.setattr('socket.create_connection',no_network)
    monkeypatch.setattr('urllib.request.urlopen',no_network)
    created=datetime.now(timezone.utc)
    output={kind:writer(content,created) for kind,writer in RENDERERS.items()}
    book,text=visible_xlsx(output['XLSX'])
    assert any(str(cell.value).startswith('=HYPERLINK') and cell.data_type=='s' for sheet in book for row in sheet for cell in row if cell.value)
    assert all(cell.data_type!='f' and not cell.hyperlink for sheet in book for row in sheet for cell in row)
    for kind in ['DOCX','XLSX']:
        with ZipFile(BytesIO(output[kind])) as z:
            assert not any('vbaProject' in name or name.endswith('.vba') for name in z.namelist())
            assert not any(b'TargetMode="External"' in z.read(name) for name in z.namelist() if name.endswith('.rels'))
    assert '{{7*7}}' in visible_docx(output['DOCX'])
    assert '<img src=' in visible_pdf(output['PDF'])


def test_m9c27_28_atomic_failure_retry_tamper_and_regeneration(proposal,monkeypatch):
    c,t,db,m,snapshot,v,comp=proposal
    first,_=generate(c,t,v,formats=['DOCX'],include_package=False);before=artifact(c,t,v,first,'DOCX')
    payload={'request_id':str(uuid4()),'formats':['PDF'],'include_package':True}
    def fail(*a,**k):raise ValueError('INJECTED_RENDER_FAILURE')
    with monkeypatch.context() as patch:
        patch.setitem(RENDERERS,'PDF',fail)
        response=c.post(route(t,v)+'/exports',json=payload)
        assert response.status_code==422 and 'INJECTED_RENDER_FAILURE' in response.text
    assert artifact(c,t,v,first,'DOCX')==before
    assert len(read(c,route(t,v)+'/exports')['items'])==1
    response=c.post(route(t,v)+'/exports',json=payload);assert response.status_code==200,response.text
    receipt=response.json();assert c.post(route(t,v)+'/exports',json=payload).json()==receipt
    assert c.post(route(t,v)+'/exports',json={**payload,'formats':['DOCX']}).status_code==409
    another,_=generate(c,t,v,formats=['PDF'])
    assert another['semantic_sha256']==receipt['semantic_sha256']
    file=_base(t,v['id'])/receipt['request_id']/'proposal.pdf';file.write_bytes(file.read_bytes()+b'tamper')
    assert c.post(route(t,v)+f"/exports/{receipt['request_id']}/verify").status_code==409
    assert c.get(route(t,v)+f"/exports/{receipt['request_id']}/files/PDF").status_code==409
    assert artifact(c,t,v,first,'DOCX')==before
    assert not list(_base(t,v['id']).glob('.tmp-*'))


def test_independent_package_tamper_and_selected_contents(proposal):
    c,t,db,m,snapshot,v,comp=proposal;receipt,_=generate(c,t,v,formats=['DOCX','XLSX'])
    original=artifact(c,t,v,receipt,'ZIP');manifest=verify_package(original)
    assert {a['type'] for a in manifest['artifacts']}=={'DOCX','XLSX'}
    edited=BytesIO()
    with ZipFile(BytesIO(original)) as source,ZipFile(edited,'w',compression=ZIP_DEFLATED) as destination:
        for name in source.namelist():destination.writestr(name,source.read(name)+(b'TAMPER' if name=='proposal.docx' else b''))
    with pytest.raises(ValueError):verify_package(edited.getvalue())


def test_m9c29_many_scope_requirement_evidence_rows_without_n_plus_one(proposal):
    from app.proposal_composition import Content,MaterializedSource,canonical_rows,_content_dump
    from app.proposal_readiness import get_snapshot,digest
    c,t,db,m,snapshot,v,comp=proposal
    original=Content.model_validate(db.get(ProposalComposition,v['id']).content)
    overlay=original.model_copy(deep=True)
    requirement_key=next(key for key in original.sources if key.startswith('REQUIREMENT:'))
    evidence_key=next(key for key in original.sources if key.startswith('REQUESTED_EVIDENCE:'))
    old_req=requirement_key.split(':',1)[1];old_ev=evidence_key.split(':',1)[1]
    new_keys=[]
    def expand(projection):
        seed=deepcopy(projection['rows'][0]);states=deepcopy(projection['tracked_context'])
        for _ in range(180):
            scope=deepcopy(seed);scope.update(scope_id=str(uuid4()),quantity_id=str(uuid4()),commercial_line_id=str(uuid4()))
            scope['quantity_scope_id']=scope['scope_id'];projection['rows'].append(scope)
            req_id=str(uuid4());ev_id=str(uuid4())
            projection['requirement_representations'].append({'requirement_id':req_id,'representation_fingerprint':original.sources[requirement_key].representation_fingerprint})
            for state in states:
                if state['entity_type']=='REQUIREMENT' and state['entity_id']==old_req:
                    clone=deepcopy(state);clone.update(id=str(uuid4()),entity_id=req_id);projection['tracked_context'].append(clone)
                if state['entity_type']=='REQUESTED_EVIDENCE' and state['entity_id']==old_ev:
                    clone=deepcopy(state);clone.update(id=str(uuid4()),entity_id=ev_id,parent_requirement_id=req_id);projection['tracked_context'].append(clone)
            for old_key,new_id in [(requirement_key,req_id),(evidence_key,ev_id)]:
                kind=old_key.split(':',1)[0];key=kind+':'+new_id
                payload=original.sources[old_key].model_dump(mode='json',exclude={'content_sha256'});payload['source_id']=new_id
                overlay.sources[key]=MaterializedSource(**payload,content_sha256=digest(payload))
                overlay.responses[key]=original.responses[old_key].model_copy(deep=True);new_keys.append(key)
    large_snapshot=synthetic_snapshot(db,snapshot,expand)
    large,_=create(c,t,large_snapshot,acknowledge_stale=True)
    canonical=canonical_rows(get_snapshot(db,t,large['snapshot_id']),overlay)
    for key in new_keys:overlay.responses[key].source_sha256=canonical[key]['source_sha256'];overlay.responses[key].source_match='NEW'
    now=datetime.now(timezone.utc)
    db.add(ProposalComposition(version_id=large['id'],contract_version='proposal-composition-v1',revision=1,content=_content_dump(overlay),created_at=now,updated_at=now));db.flush()
    large=freeze(c,t,large,{'revision':1})
    engine=db.get_bind().engine
    def measure(version):
        statements=[]
        def collect(conn,cursor,statement,parameters,context,executemany):
            if statement.lstrip().upper().startswith('SELECT'):statements.append(statement)
        event.listen(engine,'before_cursor_execute',collect)
        try:receipt,_=generate(c,t,version)
        finally:event.remove(engine,'before_cursor_execute',collect)
        assert not any('FROM requirements' in s or 'FROM tender_scope' in s for s in statements)
        return receipt,len(statements)
    small,small_queries=measure(v);big,big_queries=measure(large)
    assert big_queries<=small_queries+2,{'small':small_queries,'large':big_queries}
    content=verify_package(artifact(c,t,large,big,'ZIP'))['content']
    sizes={section['id']:len(section['rows']) for section in content['sections']}
    assert sizes['technical_scope']>=180 and sizes['requirements']>=180 and sizes['requested_evidence']>=180
    Path('reports/m9_c').mkdir(parents=True,exist_ok=True)
    Path('reports/m9_c/query_scaling.json').write_text(json.dumps({'small_selects':small_queries,'large_selects':big_queries,'rows_by_section':sizes},indent=2)+'\n')


def test_legacy_frozen_missing_literal_is_not_backfilled(ready,monkeypatch):
    from app.models import ProposalVersion
    from app.proposal_composition import canonical_rows
    from app.proposal_readiness import digest,get_snapshot
    from app.proposal_versions import frozen_envelope
    c,t,db,m,snapshot=ready;v,_=create(c,t,snapshot);finish(c,t,v)
    stored=db.get(ProposalComposition,v['id']);legacy=deepcopy(stored.content);legacy.pop('sources')
    rows=canonical_rows(get_snapshot(db,t,v['snapshot_id']))
    for key,response in legacy['responses'].items():response['source_sha256']=rows[key]['source_sha256']
    stored.content=legacy;stored.revision+=1;db.flush()
    old_hash=digest({'contract_version':'proposal-composition-v1','version_id':v['id'],'revision':stored.revision,'content':legacy})
    version=db.get(ProposalVersion,v['id']);version.state='FROZEN';version.edit_revision+=1;version.frozen_at=datetime.now(timezone.utc);version.updated_at=version.frozen_at
    version.freeze_acknowledgements={'request_sha256':'a'*64,'warnings':True,'stale':False,'source_state':'CURRENT',
        'composition_contract':'proposal-composition-v1','composition_revision':stored.revision,'composition_sha256':old_hash}
    version.frozen_sha256=digest(frozen_envelope(version));db.flush();before=version.frozen_sha256
    monkeypatch.setattr('app.requirement_matrix.get_tender_requirement_matrix',lambda *a:pytest.fail('Legacy mutable backfill'))
    response=c.post(route(t,v)+'/exports',json={'request_id':str(uuid4())})
    assert response.status_code==409 and 'SOURCE_LITERAL_UNAVAILABLE' in response.text
    assert version.frozen_sha256==before and 'sources' not in stored.content


def test_excel_numeric_precision_and_negative_values_are_preserved():
    from openpyxl import Workbook
    from app.proposal_export_renderers import _number
    from decimal import Decimal
    sheet=Workbook().active
    for index,value in enumerate(['-12.50','12345678901234567.89','0.000000000000000001','-0.00'],1):
        cell=sheet.cell(index,1);_number(cell,value)
        assert Decimal(str(cell.value))==Decimal(value)
        assert cell.data_type!='f'
    assert sheet.cell(1,1).data_type=='n' and sheet.cell(2,1).data_type=='s'
