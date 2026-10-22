"""M7C-01–30 bounded synthetic round trips through canonical APIs."""
import csv
import io
import json
from decimal import Decimal
from uuid import uuid4
from zipfile import ZipFile
import pytest
from openpyxl import load_workbook
from sqlalchemy import select
from app.models import CommercialLine, CommercialLineRevision
from app import commercial_workbook as w
from tests.test_commercial_matrix import sheet, matrix, bulk
from tests.test_workspace_flow import scene


def export(c,r,kind='xlsx'):
    response=c.get(r.replace('/commercial-matrix',f'/commercial-exchange/export/{kind}'))
    assert response.status_code==200,response.text
    return response.content


def edit(data, updates, *, reverse=False, remove=False):
    book=load_workbook(io.BytesIO(data));ws=book[w.SHEET]
    headers={cell.value:cell.column for cell in ws[1]}
    for (row,key),value in updates.items():ws.cell(row,headers[key]).value=value
    if reverse:
        a=[cell.value for cell in ws[2]];b=[cell.value for cell in ws[3]]
        for col,(x,y) in enumerate(zip(a,b),1):ws.cell(2,col).value=y;ws.cell(3,col).value=x
    if remove:ws.delete_rows(3)
    out=io.BytesIO();book.save(out);book.close();return out.getvalue()


def stage(c,r,data,kind='xlsx',**flags):
    response=c.post(r.replace('/commercial-matrix','/commercial-exchange/stage'),files={'file':(f'precios.{kind}',data)},data={k:str(v).lower() for k,v in flags.items()})
    assert response.status_code==200,response.text
    return response.json()


def apply(c,r,preview):
    response=c.post(r.replace('/commercial-matrix',f"/commercial-exchange/stages/{preview['stage_id']}/apply"),json={'confirm':True})
    assert response.status_code==200,response.text
    return response.json()


def test_m7c01_04_export_identity_sort_and_preview_pure(sheet):
    c,_,db,r=sheet;before=matrix(c,r)
    data=export(c,r);book=load_workbook(io.BytesIO(data));ws=book[w.SHEET]
    assert ws.max_row==3 and ws.cell(2,1).value==before['rows'][0]['scope_id']
    assert book[w.META].sheet_state=='hidden' and not ws.protection.sheet
    book.close()
    edited=edit(data,{(2,'unit_cost'):'100.01',(2,'currency'):'MXN'},reverse=True)
    preview=stage(c,r,edited)
    changed=next(row for row in preview['rows'] if row['state']=='CHANGED')
    assert changed['scope_id']==before['rows'][0]['scope_id'] and changed['eligible']
    assert changed['diff'] and not list(db.scalars(select(CommercialLine)))
    apply(c,r,preview)
    assert Decimal(matrix(c,r)['rows'][0]['inputs']['unit_cost'])==Decimal('100.01')


@pytest.mark.parametrize('field,value',[('description','alterado'),('quantity','999'),('unit','KG'),('commercial_line_id',str(uuid4())),('source_document',str(uuid4()))])
def test_m7c06_07_reference_tamper_never_writes(sheet,field,value):
    c,t,db,r=sheet;before=c.get(f'/tenders/{t}/workspace').json()['scopes']
    data=edit(export(c,r),{(2,field):value,(2,'unit_cost'):'100'})
    preview=stage(c,r,data)
    assert preview['rows'][0]['errors']==['REFERENCE_FIELDS_MODIFIED'] and not preview['rows'][0]['eligible']
    apply(c,r,preview)
    assert not list(db.scalars(select(CommercialLine))) and c.get(f'/tenders/{t}/workspace').json()['scopes']==before


@pytest.mark.parametrize('field,value,state',[('scope_id',str(uuid4()),'UNKNOWN_SCOPE'),('scope_id','../x','INVALID'),('currency','ZZZ','INVALID'),('unit_cost','1,234.56','INVALID'),('unit_cost','NaN','INVALID')])
def test_m7c08_10_invalid_unknown(sheet,field,value,state):
    c,_,db,r=sheet
    preview=stage(c,r,edit(export(c,r),{(2,field):value}))
    assert preview['rows'][0]['state']==state
    apply(c,r,preview);assert not list(db.scalars(select(CommercialLine)))


def test_m7c09_duplicates_all_rejected(sheet):
    c,_,_,r=sheet;data=export(c,r);book=load_workbook(io.BytesIO(data));identity=book[w.SHEET].cell(2,1).value;book.close()
    preview=stage(c,r,edit(data,{(3,'scope_id'):identity,(2,'unit_cost'):'100'}))
    assert preview['counts']['DUPLICATE_SCOPE']==2


@pytest.mark.parametrize('basis,field,value,number_format,expected',[('MARKUP','markup_percent',25,'0','125'),('MARKUP','markup_percent',0.25,'0.00%','125'),('MARKUP','markup_percent','25%','@','125'),('MARGIN','target_margin_percent',0.20,'0%','125')])
def test_m7c11_13_numeric_percent_cost_first(sheet,basis,field,value,number_format,expected):
    c,_,_,r=sheet
    data=edit(export(c,r),{(2,'unit_cost'):100,(2,'currency'):'MXN',(2,'calculation_mode'):'COST_FIRST',(2,'calculation_basis'):basis,(2,field):value})
    book=load_workbook(io.BytesIO(data));ws=book[w.SHEET];col=[x.value for x in ws[1]].index(field)+1;ws.cell(2,col).number_format=number_format
    out=io.BytesIO();book.save(out);book.close()
    preview=stage(c,r,out.getvalue());assert preview['rows'][0]['eligible']
    apply(c,r,preview);row=matrix(c,r)['rows'][0]
    assert Decimal(row['inputs']['offered_unit_price'])==Decimal(expected) and row['price_method']=='CALCULATED'


def test_m7c14_16_explicit_price_and_backend_derived(sheet):
    c,_,_,r=sheet
    data=edit(export(c,r),{(2,'unit_cost'):'100',(2,'currency'):'MXN',(2,'offered_unit_price'):'90',(2,'offered_total'):'999999',(2,'notes'):'Árbol — precio humano'})
    preview=stage(c,r,data);assert 'DERIVED_VALUES_IGNORED' in preview['rows'][0]['warnings']
    apply(c,r,preview);row=matrix(c,r)['rows'][0]
    assert row['price_method']=='MANUAL' and Decimal(row['inputs']['offered_unit_price'])==90
    assert row['summary']['offered_total']!='999999' and 'PRICE_BELOW_COST' in row['summary']['warnings']


def test_m7c15_20_manual_protection_and_idempotence(sheet):
    c,_,db,r=sheet
    bulk(c,r,matrix(c,r)['rows'],changes={'unit_cost':'100','currency':'MXN','offered_unit_price':'150'},apply=True)
    data=edit(export(c,r),{(2,'calculation_mode'):'COST_FIRST',(2,'calculation_basis'):'MARKUP',(2,'markup_percent'):'25'})
    preview=stage(c,r,data);assert preview['rows'][0]['price_outcome']=='PROTECTED'
    apply(c,r,preview);n=len(list(db.scalars(select(CommercialLineRevision))))
    assert Decimal(matrix(c,r)['rows'][0]['inputs']['offered_unit_price'])==150
    apply(c,r,preview);assert len(list(db.scalars(select(CommercialLineRevision))))==n
    again=stage(c,r,data);assert again['rows'][0]['state']=='UNCHANGED'
    apply(c,r,again);assert len(list(db.scalars(select(CommercialLineRevision))))==n


def test_m7c21_export_and_preview_stale(sheet):
    c,_,_,r=sheet
    data=edit(export(c,r),{(2,'unit_cost'):'100',(2,'currency'):'MXN'})
    preview=stage(c,r,data)
    bulk(c,r,[matrix(c,r)['rows'][0]],changes={'unit_cost':'200'},apply=True)
    result=apply(c,r,preview)
    assert result['counts']=={'CONFLICT':1}
    again=stage(c,r,data);assert again['rows'][0]['state']=='CONFLICT'
    assert Decimal(matrix(c,r)['rows'][0]['inputs']['unit_cost'])==200


def test_m7c22_missing_no_delete(sheet):
    c,_,db,r=sheet
    bulk(c,r,matrix(c,r)['rows'],changes={'currency':'MXN'},apply=True)
    data=edit(export(c,r),{(2,'unit_cost'):'100'},remove=True)
    preview=stage(c,r,data);assert preview['counts']['MISSING_FROM_IMPORT']==1
    apply(c,r,preview);assert len(list(db.scalars(select(CommercialLine))))==2 and len(matrix(c,r)['rows'])==2


def test_m7c24_25_27_28_csv_utf8_hash_and_origin(sheet):
    c,_,db,r=sheet;data=export(c,r,'csv');records=list(csv.reader(io.StringIO(data.decode('utf-8'))))
    headers=records[1];records[2][headers.index('notes')]='Árbol, válvula — México';records[2][headers.index('unit_cost')]='0.10';records[2][headers.index('currency')]='USD'
    out=io.StringIO(newline='');csv.writer(out).writerows(records);raw=out.getvalue().encode('utf-8')
    preview=stage(c,r,raw,'csv');apply(c,r,preview)
    row=matrix(c,r)['rows'][0];assert row['inputs']['notes']=='Árbol, válvula — México'
    revision=db.scalar(select(CommercialLineRevision))
    assert revision.action=='CSV_IMPORT' and revision.input_origin=='CSV_IMPORT'
    assert revision.calculation_snapshot['bulk_operation']['import_evidence']['sha256']==preview['file_evidence']['sha256']


@pytest.mark.parametrize('field,valid',[('unit_cost',False),('notes',False),('offered_total',True)])
def test_m7c26_formulas_never_evaluated(sheet,field,valid):
    c,_,_,r=sheet
    data=edit(export(c,r),{(2,field):'=HYPERLINK("https://invalid.test","1")',(2,'currency'):'MXN'})
    preview=stage(c,r,data)
    assert preview['rows'][0]['eligible']==valid
    if valid:assert 'DERIVED_VALUES_IGNORED' in preview['rows'][0]['warnings']


def test_metadata_version_and_tender_fail_closed(sheet):
    c,_,_,r=sheet
    book=load_workbook(io.BytesIO(export(c,r)));book[w.META].cell(1,2,'unknown-version');out=io.BytesIO();book.save(out);book.close()
    response=c.post(r.replace('/commercial-matrix','/commercial-exchange/stage'),files={'file':('x.xlsx',out.getvalue())})
    assert response.status_code==422


def test_m7c19_30_valid_apply_keeps_mvp06_and_requires_confirmation(sheet):
    c,t,db,r=sheet;before=c.get(f'/tenders/{t}/workspace').json()['scopes']
    preview=stage(c,r,edit(export(c,r),{(2,'unit_cost'):'100',(2,'currency'):'MXN'}))
    route=r.replace('/commercial-matrix',f"/commercial-exchange/stages/{preview['stage_id']}/apply")
    assert c.post(route,json={'confirm':False}).status_code==422
    assert not list(db.scalars(select(CommercialLine)))
    apply(c,r,preview)
    assert c.get(f'/tenders/{t}/workspace').json()['scopes']==before
    assert db.scalar(select(CommercialLineRevision)).action=='XLSX_IMPORT'


def rewrite_zip(data, changes):
    output=io.BytesIO()
    with ZipFile(io.BytesIO(data)) as source, ZipFile(output,'w') as destination:
        for info in source.infolist():destination.writestr(info.filename,changes.pop(info.filename,source.read(info.filename)))
        for name,value in changes.items():destination.writestr(name,value)
    return output.getvalue()


def test_lexical_numeric_precision_does_not_round_through_float(sheet):
    from openpyxl.utils import get_column_letter
    import xml.etree.ElementTree as ET
    c,_,_,r=sheet
    data=edit(export(c,r),{(2,'unit_cost'):1,(2,'currency'):'MXN'})
    coord=f"{get_column_letter(w.HEADERS.index('unit_cost')+1)}2"
    with ZipFile(io.BytesIO(data)) as archive:node=ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
    cell=next(n for n in node.iter(f"{{{w.NS['s']}}}c") if n.attrib['r']==coord)
    exact='90071992547409.123456';cell.find('s:v',w.NS).text=exact
    data=rewrite_zip(data,{'xl/worksheets/sheet1.xml':ET.tostring(node)})
    preview=stage(c,r,data);apply(c,r,preview)
    assert Decimal(matrix(c,r)['rows'][0]['inputs']['unit_cost'])==Decimal(exact)


@pytest.mark.parametrize('cell_value',['=SUM(1,2)','+CMD','@SUM(1,2)','-command'])
def test_export_text_formula_injection_is_literal_in_both_formats(sheet,cell_value):
    c,_,_,r=sheet
    bulk(c,r,[matrix(c,r)['rows'][0]],changes={'notes':cell_value},apply=True)
    data=export(c,r);book=load_workbook(io.BytesIO(data));cell=book[w.SHEET].cell(2,w.HEADERS.index('notes')+1)
    assert cell.data_type=='s' and cell.value==cell_value;book.close()
    records=list(csv.reader(io.StringIO(export(c,r,'csv').decode())))
    assert records[2][records[1].index('notes')]=="'"+cell_value
    assert stage(c,r,export(c,r,'csv'),'csv')['rows'][0]['state']=='UNCHANGED'


@pytest.mark.parametrize('value',[True,'',None])
def test_invalid_boolean_or_empty_value_and_unchanged_blanks(sheet,value):
    c,_,_,r=sheet
    if value is True:
        preview=stage(c,r,edit(export(c,r),{(2,'unit_cost'):value}))
        assert preview['rows'][0]['state']=='INVALID'
    else:
        bulk(c,r,[matrix(c,r)['rows'][0]],changes={'notes':value},apply=True)
        preview=stage(c,r,export(c,r))
        assert preview['rows'][0]['state']=='UNCHANGED'


def test_explicit_imported_price_wins_cost_first_and_currency_requires_permission(sheet):
    c,_,_,r=sheet
    bulk(c,r,[matrix(c,r)['rows'][0]],changes={'unit_cost':'100','currency':'MXN','offered_unit_price':'150'},apply=True)
    data=edit(export(c,r),{(2,'currency'):'USD',(2,'offered_unit_price'):'170',(2,'calculation_mode'):'COST_FIRST',(2,'calculation_basis'):'MARKUP',(2,'markup_percent'):'25'})
    denied=stage(c,r,data);assert denied['rows'][0]['state']=='CONFLICT' and 'CURRENCY_CONFIRMATION_REQUIRED' in denied['rows'][0]['errors']
    allowed=stage(c,r,data,confirm_currency_change=True);apply(c,r,allowed)
    row=matrix(c,r)['rows'][0]
    assert Decimal(row['inputs']['offered_unit_price'])==170 and row['inputs']['currency']=='USD' and row['price_method']=='MANUAL'


def test_explicit_recalculate_manual_import_option(sheet):
    c,_,_,r=sheet
    bulk(c,r,[matrix(c,r)['rows'][0]],changes={'unit_cost':'100','currency':'MXN','offered_unit_price':'150'},apply=True)
    data=edit(export(c,r),{(2,'calculation_mode'):'COST_FIRST',(2,'calculation_basis'):'MARKUP',(2,'markup_percent'):'25'})
    preview=stage(c,r,data,replace_manual_price=True);apply(c,r,preview)
    assert Decimal(matrix(c,r)['rows'][0]['inputs']['offered_unit_price'])==125


@pytest.mark.parametrize('name,body',[('xl/vbaProject.bin',b'not executed'),('../escape.xml',b'<xml/>'),('xl/externalLinks/externalLink1.xml',b'<xml/>'),('evil.xml',b'<!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>'),('unsafe.rels',b'<Relationships><Relationship TargetMode="External" Target="https://invalid.test"/></Relationships>')])
def test_archive_untrusted_parts_fail_before_loading(name,body):
    data=w.write_workbook([],{'format_version':w.VERSION,'tender_id':str(uuid4()),'export_id':str(uuid4()),'export_timestamp':'x'},'xlsx')
    from fastapi import HTTPException
    with pytest.raises(HTTPException):w.read_workbook(rewrite_zip(data,{name:body}),'xlsx')


def test_file_rows_and_format_bounds():
    from fastapi import HTTPException
    with pytest.raises(HTTPException):w.read_workbook(b'x'*(w.MAX_FILE+1),'xlsx')
    with pytest.raises(HTTPException):w.read_workbook(b'not a workbook','xlsx')
    with pytest.raises(HTTPException):w.read_workbook(b'\xffbad','csv')
    with pytest.raises(HTTPException):w.write_workbook([{}]*(w.MAX_ROWS+1),{},'csv')


def test_context_expiry_corruption_and_tender_binding(sheet):
    from app.commercial_exchange_store import directory, canonical
    import hashlib
    c,t,_,r=sheet
    data=edit(export(c,r),{(2,'unit_cost'):'100'})
    preview=stage(c,r,data)
    path=directory()/f"stage_{preview['stage_id']}.json"
    assert path.stat().st_mode & 0o777==0o600 and directory().stat().st_mode & 0o777==0o700
    body=json.loads(path.read_text());body['payload']['expires_at']='2000-01-01T00:00:00+00:00';body['sha256']=hashlib.sha256(canonical(body['payload'])).hexdigest();path.write_bytes(canonical(body))
    endpoint=r.replace('/commercial-matrix',f"/commercial-exchange/stages/{preview['stage_id']}/apply")
    assert c.post(endpoint,json={'confirm':True}).status_code==410
    path.write_text('{}')
    assert c.post(endpoint,json={'confirm':True}).status_code==409
    path.unlink()  # Remove only this deliberately corrupted test fixture.
    # Wrong Tender cannot read a valid export context.
    preview=stage(c,r,data)
    other=c.post('/tenders',json={'title':'Other'}).json()['id']
    assert c.post(f"/tenders/{other}/commercial-exchange/stages/{preview['stage_id']}/apply",json={'confirm':True}).status_code==409


def test_reserved_import_origin_and_upload_transport_limit(sheet):
    c,_,_,r=sheet
    rows=matrix(c,r)['rows']
    from tests.test_commercial_matrix import target
    assert c.post(r+'/bulk',json={'targets':[target(rows[0])],'changes':{'unit_cost':'1'},'input_origin':'XLSX_IMPORT'}).status_code==422
    response=c.post(r.replace('/commercial-matrix','/commercial-exchange/stage'),content=b'bad',headers={'Content-Length':str(w.MAX_FILE+65537)})
    assert response.status_code==413
