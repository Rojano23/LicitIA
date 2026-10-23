"""Neutral local renderers over an allowlisted export DTO; no SQL, URLs or templates."""
from io import BytesIO
from decimal import Decimal,InvalidOperation
from pathlib import Path
from html import escape
import json
from threading import RLock
PDF_LOCK=RLock()
from app.proposal_export_contract import artifact_identity

SHEETS={'identification':'Resumen','technical_scope':'Alcance Tecnico','requirements':'Requisitos',
    'requested_evidence':'Evidencias','economic':'Propuesta Economica','clarifications':'Aclaraciones',
    'annex_references':'Referencias','notes':'Notas'}


def docx_bytes(content,created_at):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Pt
    document=Document();style=document.styles['Normal'];style.font.name='Arial';style.font.size=Pt(10)
    document.sections[0].header.paragraphs[0].text=content.proposal_state
    document.add_heading(content.proposal_title,0)
    document.add_paragraph(content.notice)
    if content.warnings_acknowledged:document.add_paragraph('La propuesta se generó con advertencias reconocidas por el participante.')
    props=document.core_properties;props.title=content.proposal_title[:255];props.author=(content.bidder_name or '')[:255]
    props.identifier=content.version_id;props.subject='snapshot:'+content.snapshot_id
    props.keywords='proposal-export-v1;artifact:DOCX;state:'+content.proposal_state
    props.comments='semantic:'+content.semantic_sha256+';composition:'+content.composition_sha256
    props.created=created_at.replace(tzinfo=None);props.modified=props.created
    for section in content.sections:
        document.add_heading(section.title,1)
        groups=list(dict.fromkeys(row.group for row in section.rows)) or [None]
        for group in groups:
            if group:document.add_heading('Moneda: '+group,2)
            table=document.add_table(rows=1,cols=len(section.columns));table.style='Table Grid'
            marker=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(marker)
            for cell,column in zip(table.rows[0].cells,section.columns):cell.text=column.label
            for row in section.rows:
                if row.group!=group:continue
                for cell,column in zip(table.add_row().cells,section.columns):cell.text=row.values.get(column.key,'')
        if not section.rows:document.add_paragraph('Sin filas incluidas.')
    out=BytesIO();document.save(out);return out.getvalue()


def _text(cell,value):
    if len(value)>32767:raise ValueError('EXPORT_CELL_TEXT_TOO_LONG')
    cell.value=value;cell.data_type='s'
    if value.lstrip().startswith(('=','+','-','@')):cell.quotePrefix=True


def _number(cell,value):
    try:number=Decimal(value)
    except InvalidOperation as exc:raise ValueError('INVALID_CANONICAL_EXPORT_NUMBER') from exc
    if not number.is_finite():raise ValueError('INVALID_CANONICAL_EXPORT_NUMBER')
    exponent=number.as_tuple().exponent
    if len(number.as_tuple().digits)<=15 and -12<=exponent<=0 and not(number.is_zero() and number.is_signed()):
        numeric=float(number)
        if Decimal(str(numeric))==number:
            cell.value=numeric;cell.number_format='0'+('.'+'0'*(-exponent) if exponent<0 else '');return
    _text(cell,value)


def xlsx_bytes(content,created_at):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font
    from openpyxl.utils import get_column_letter
    book=Workbook();summary=book.active;summary.title='Resumen'
    _text(summary.cell(1,1),'Propuesta · '+content.proposal_state)
    _text(summary.cell(2,1),content.notice)
    book.properties.title=content.proposal_title;book.properties.identifier=content.version_id
    book.properties.subject='snapshot:'+content.snapshot_id
    book.properties.description=json.dumps(artifact_identity(content,created_at,'XLSX'),ensure_ascii=False,sort_keys=True)
    book.properties.created=created_at.replace(tzinfo=None);book.properties.modified=book.properties.created
    metadata=book.create_sheet('_Metadata');metadata.sheet_state='hidden'
    meta=artifact_identity(content,created_at,'XLSX')
    for n,(key,value) in enumerate(meta.items(),1):_text(metadata.cell(n,1),key);_text(metadata.cell(n,2),str(value))
    mapping_start=len(meta)+2
    for section in content.sections:
        sheet=summary if section.id=='identification' else book.create_sheet(SHEETS[section.id])
        _text(sheet.cell(1,1),section.title+' · '+content.proposal_state)
        _text(sheet.cell(2,1),content.notice)
        for column_index,column in enumerate(section.columns,1):
            cell=sheet.cell(3,column_index);_text(cell,column.label);cell.font=Font(bold=True)
            sheet.column_dimensions[get_column_letter(column_index)].width=45 if column.key in {'description','text','response'} else 24
        for row_index,row in enumerate(section.rows,4):
            for column_index,column in enumerate(section.columns,1):
                cell=sheet.cell(row_index,column_index);value=row.values.get(column.key,'')
                if column.numeric and value:_number(cell,value)
                else:_text(cell,value)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            for j,value in enumerate([section.id,str(row_index),row.key,row.document_id or ''],1):_text(metadata.cell(mapping_start,j),value)
            mapping_start+=1
        sheet.freeze_panes='A4';sheet.auto_filter.ref=f'A3:{get_column_letter(len(section.columns))}{max(3,3+len(section.rows))}'
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
    out=BytesIO();book.save(out);return out.getvalue()


def _pdf_bytes(content,created_at):
    import reportlab
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.pagesizes import A4
    font='ProposalLocalVera'
    if font not in pdfmetrics.getRegisteredFontNames():
        location=Path(reportlab.__file__).parent/'fonts/Vera.ttf'
        if not location.is_file():raise ValueError('LOCAL_EXPORT_FONT_UNAVAILABLE')
        pdfmetrics.registerFont(TTFont(font,str(location)))
    glyphs=pdfmetrics.getFont(font).face.charToGlyph
    text=''.join([content.proposal_title,content.notice,*[s.title+''.join(c.label for c in s.columns)+''.join(v for r in s.rows for v in r.values.values()) for s in content.sections]])
    if any(ord(c)>32 and ord(c) not in glyphs for c in text):raise ValueError('UNSUPPORTED_PDF_GLYPH')
    normal=ParagraphStyle('ProposalNormal',fontName=font,fontSize=9,leading=12,spaceAfter=5)
    heading=ParagraphStyle('ProposalHeading',parent=normal,fontSize=13,leading=17,spaceBefore=10,spaceAfter=8,keepWithNext=True)
    title=ParagraphStyle('ProposalTitle',parent=heading,fontSize=17,leading=21)
    def paragraph(value,style=normal):return Paragraph(escape(value).replace('\n','<br/>'),style)
    story=[paragraph(content.proposal_title,title),paragraph(content.proposal_state),paragraph(content.notice)]
    if content.warnings_acknowledged:story.append(paragraph('La propuesta se generó con advertencias reconocidas por el participante.'))
    for section in content.sections:
        story.append(paragraph(section.title,heading));current=None
        for row in section.rows:
            if row.group and row.group!=current:story.append(paragraph('Moneda: '+row.group,heading));current=row.group
            for column in section.columns:
                story.append(paragraph(column.label+': '+row.values.get(column.key,'')))
            story.append(Spacer(1,7))
        if not section.rows:story.append(paragraph('Sin filas incluidas.'))
    identity=json.dumps(artifact_identity(content,created_at,'PDF'),ensure_ascii=False,sort_keys=True)
    def page(canvas,document):
        canvas.setTitle(content.proposal_title);canvas.setAuthor(content.bidder_name or '')
        canvas.setSubject(identity);canvas.setCreator('proposal-export-v1')
        canvas.setFont(font,8);canvas.drawString(36,20,content.proposal_state)
        canvas.drawRightString(A4[0]-36,20,'Página '+str(canvas.getPageNumber()))
    out=BytesIO();document=SimpleDocTemplate(out,pagesize=A4,leftMargin=36,rightMargin=36,topMargin=36,bottomMargin=36,invariant=1)
    document.build(story,onFirstPage=page,onLaterPages=page)
    return out.getvalue()


def pdf_bytes(content,created_at):
    with PDF_LOCK:return _pdf_bytes(content,created_at)


RENDERERS={'DOCX':docx_bytes,'XLSX':xlsx_bytes,'PDF':pdf_bytes}
