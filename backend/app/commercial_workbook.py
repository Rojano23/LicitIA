"""Versioned, bounded XLSX/CSV transport. No formulas are calculated."""
import csv
import io
import json
import posixpath
import re
from decimal import Decimal, InvalidOperation, localcontext, Context
from pathlib import PurePosixPath
from zipfile import ZipFile, BadZipFile
import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Protection, Font
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from fastapi import HTTPException

VERSION = 'licitia-commercial-workbook-v1'
MAX_FILE = 5 * 1024 * 1024
MAX_EXPANDED = 64 * 1024 * 1024
MAX_ROWS = 1000
MAX_TEXT = 32767
SHEET = 'Commercial Matrix'
META = 'LicitIA Metadata'
REFERENCES = ('scope_id', 'commercial_line_id', 'revision', 'source_fingerprint', 'item_reference',
              'description', 'quantity', 'unit', 'source_document', 'source_page', 'source_locator')
EDITABLE = ('cost_source_type', 'cost_value_origin', 'reference_party_name', 'reference_party_reference',
            'reference_party_note', 'unit_cost', 'currency', 'calculation_mode', 'calculation_basis',
            'markup_percent', 'target_margin_percent', 'offered_unit_price', 'notes')
DERIVED = ('total_cost', 'offered_total', 'profit', 'actual_markup_percent', 'actual_margin_percent', 'commercial_status', 'warnings')
HEADERS = REFERENCES + EDITABLE + DERIVED
NUMERIC = {'unit_cost', 'offered_unit_price', 'markup_percent', 'target_margin_percent'}
PERCENTS = {'markup_percent', 'target_margin_percent'}
NUMBER = re.compile(r'^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$')
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def fail(code, status=422):
    raise HTTPException(status, code)


def text(value):
    if value is None:
        return ''
    return ''.join(c for c in str(value) if ord(c) >= 32 or c in '\t\r\n')[:MAX_TEXT]


def reference_row(row):
    source = row['source']['scope']['source'] if row['source']['scope'] else {}
    quantities = row['source']['quantities']
    selected = next((q for q in quantities if q['id'] == row['source']['selected_quantity_id']), None)
    return {k: text(v) for k, v in {
        'scope_id': row['scope_id'], 'commercial_line_id': row['commercial_line_id'], 'revision': row['revision'],
        'source_fingerprint': row['source_fingerprint'], 'item_reference': row['item_reference'], 'description': row['description'],
        'quantity': selected['quantity_raw'] if selected else ' | '.join(q['quantity_raw'] for q in quantities),
        'unit': selected['unit_raw'] if selected else ' | '.join(q['unit_raw'] or '' for q in quantities),
        'source_document': source.get('document_id'), 'source_page': source.get('page'), 'source_locator': source.get('locator'),
    }.items()}


def export_row(row):
    fields = {k: '' if row['inputs'][k] is None else str(row['inputs'][k]) for k in EDITABLE}
    if any(len(v) > MAX_TEXT or any(ord(c) < 32 and c not in '\t\r\n' or ord(c) in {0xFFFE, 0xFFFF} for c in v) for v in fields.values()):
        fail('UNREPRESENTABLE_COMMERCIAL_EXPORT_TEXT')
    return reference_row(row) | fields | {
        k: text(row['summary'].get(k)) for k in DERIVED if k not in {'commercial_status', 'warnings'}
    } | {'commercial_status': row['summary']['status'], 'warnings': ', '.join(row['summary']['warnings'])}


def _csv_escape(value, header):
    if value and value[0] in '=+-@\t\r' and not (header in NUMERIC and NUMBER.fullmatch(value)):
        return "'" + value
    return value


def write_workbook(rows, metadata, kind):
    if len(rows) > MAX_ROWS:
        fail('SPREADSHEET_ROW_LIMIT', 413)
    if kind == 'csv':
        output = io.StringIO(newline='')
        writer = csv.writer(output, delimiter=',', lineterminator='\r\n')
        writer.writerow([metadata[k] for k in ['format_version', 'tender_id', 'export_id', 'export_timestamp']])
        writer.writerow(HEADERS)
        for row in rows:
            writer.writerow([_csv_escape(row[k], k) for k in HEADERS])
        result = output.getvalue().encode('utf-8')
    elif kind == 'xlsx':
        book = Workbook()
        sheet = book.active; sheet.title = SHEET
        sheet.append(HEADERS)
        for column in sheet[1]:
            column.font = Font(bold=True)
        for index, row in enumerate(rows, 2):
            for col, key in enumerate(HEADERS, 1):
                cell = sheet.cell(index, col, row[key])
                cell.data_type = 's'  # Even notes beginning '=' are literal strings.
                cell.number_format = '@'
                cell.protection = Protection(locked=key not in EDITABLE)
                cell.fill = PatternFill('solid', fgColor='DBEAFE' if key in EDITABLE else 'E5E7EB')
        sheet.freeze_panes = 'F2'; sheet.auto_filter.ref = sheet.dimensions
        sheet.protection.sheet = False  # Keep sorting/moving full rows possible.
        instructions = book.create_sheet('Instructions')
        for line in [VERSION, 'Azul: editable comercial. Gris: referencia; nunca modifica evidencia.',
            'Ordene filas completas conservando Scope ID. No cambie headers ni metadata.',
            'Pegue valores; fórmulas editables se rechazan. Backend recalcula totales.',
            'Números: punto decimal, sin miles. Markup/margen: 25 o 25% significa 25%.',
            'Precio editado explícito prevalece. Precio sin editar conserva su método/protección.',
            'Proteja este archivo local: contiene información comercial confidencial.',
            'Límite 1000 filas/5 MiB. Contexto local válido 7 días; no es una propuesta.']:
            instructions.append([line])
        meta = book.create_sheet(META)
        for key, value in metadata.items():
            meta.append([key, value])
        meta.sheet_state = 'hidden'; meta.protection.sheet = True
        output = io.BytesIO(); book.save(output); book.close(); result = output.getvalue()
        inspect_archive(result)  # Export must satisfy the same bounded archive contract.
    else:
        fail('UNSUPPORTED_SPREADSHEET_FORMAT')
    if len(result) > MAX_FILE:
        fail('SPREADSHEET_FILE_LIMIT', 413)
    return result


def _xml(raw):
    try:
        decoded = raw.decode('utf-8-sig')
        encoding = re.search(r'<\?xml[^>]*encoding=[\'\"]([^\'\"]+)', decoded)
        if '\x00' in decoded or (encoding and encoding[1].lower() not in {'utf-8', 'utf8', 'us-ascii'}) or '<!DOCTYPE' in decoded.upper() or '<!ENTITY' in decoded.upper():
            fail('UNSAFE_SPREADSHEET_XML')
        return ET.fromstring(decoded)
    except (UnicodeError, ET.ParseError, ValueError) as exc:
        raise HTTPException(422, 'INVALID_SPREADSHEET_XML') from exc


def inspect_archive(data):
    if len(data) > MAX_FILE:
        fail('SPREADSHEET_FILE_LIMIT', 413)
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 64 or len({i.filename for i in infos}) != len(infos) or sum(i.file_size for i in infos) > MAX_EXPANDED:
                fail('SPREADSHEET_ARCHIVE_LIMIT', 413)
            xml = {}
            for info in infos:
                name = info.filename
                parts = PurePosixPath(name).parts
                if name.startswith('/') or '..' in parts or '\\' in name or info.flag_bits & 1 or info.is_dir() or 'vba' in name.lower() or 'externallink' in name.lower() or not name.endswith(('.xml', '.rels')):
                    fail('UNSUPPORTED_SPREADSHEET_ARCHIVE_CONTENT')
                raw = archive.read(info)
                node = _xml(raw)
                xml[name] = node
                if any(n.attrib.get('TargetMode', '').lower() == 'external' or 'macroEnabled' in n.attrib.get('ContentType', '') for n in node.iter()):
                    fail('EXTERNAL_OR_MACRO_SPREADSHEET_CONTENT')
                if name.startswith('xl/worksheets/'):
                    cells = list(node.iter(f"{{{NS['s']}}}c"))
                    if len(cells) > (MAX_ROWS + 1) * len(HEADERS):
                        fail('SPREADSHEET_CELL_LIMIT', 413)
                    for c in cells:
                        row, col = coordinate_to_tuple(c.attrib['r'])
                        if row > MAX_ROWS + 1 or col > len(HEADERS):
                            fail('SPREADSHEET_DIMENSION_LIMIT', 413)
            return xml
    except (BadZipFile, KeyError, OSError, RuntimeError, ValueError) as exc:
        raise HTTPException(422, 'INVALID_XLSX_ARCHIVE') from exc


def _headers(values):
    if len(values) != len(HEADERS) or len(set(values)) != len(values) or set(values) != set(HEADERS):
        fail('INVALID_SPREADSHEET_HEADERS')


def read_workbook(data, kind):
    if not data or len(data) > MAX_FILE:
        fail('SPREADSHEET_FILE_LIMIT', 413)
    if kind == 'csv':
        try:
            reader = csv.reader(io.StringIO(data.decode('utf-8-sig'), newline=''), delimiter=',', strict=True)
            first = next(reader); headers = next(reader)
            if len(first) != 4:
                fail('INVALID_CSV_METADATA')
            metadata = dict(zip(['format_version', 'tender_id', 'export_id', 'export_timestamp'], first))
            _headers(headers)
            rows = []
            for index, fields in enumerate(reader, 3):
                if index > MAX_ROWS + 2:
                    fail('SPREADSHEET_ROW_LIMIT', 413)
                if not fields:
                    continue
                if len(fields) != len(headers) or any(len(v) > MAX_TEXT for v in fields):
                    fail('INVALID_CSV_ROW')
                rows.append({k: {'value': v[1:] if len(v) > 1 and v[0] == "'" and v[1] in '=+-@\t\r' else v, 'type': 's', 'format': ''} for k, v in zip(headers, fields)})
            return metadata, rows
        except (UnicodeError, csv.Error, StopIteration) as exc:
            raise HTTPException(422, 'INVALID_UTF8_CSV') from exc
    if kind != 'xlsx':
        fail('UNSUPPORTED_SPREADSHEET_FORMAT')
    xml = inspect_archive(data)
    book = None
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False, keep_vba=False)
        if set(book.sheetnames) != {SHEET, 'Instructions', META}:
            fail('INVALID_SPREADSHEET_SHEETS')
        metadata = {}
        for record in book[META].iter_rows(max_row=8, max_col=2):
            key, value = record
            if key.value is not None:
                if key.data_type == 'f' or value.data_type == 'f' or not isinstance(key.value, str) or key.value in metadata:
                    fail('INVALID_SPREADSHEET_METADATA')
                metadata[key.value] = str(value.value)
        # Map sheet by workbook relationship, not guessed sheet number.
        workbook = xml['xl/workbook.xml']
        rid = next(n.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'] for n in workbook.find('s:sheets', NS) if n.attrib['name'] == SHEET)
        relation = next(n for n in xml['xl/_rels/workbook.xml.rels'] if n.attrib['Id'] == rid)
        path = relation.attrib['Target'].lstrip('/') if relation.attrib['Target'].startswith('/') else posixpath.normpath('xl/' + relation.attrib['Target'])
        sheet_xml = xml[path]
        raw_cells = {n.attrib['r']: n for n in sheet_xml.iter(f"{{{NS['s']}}}c")}
        max_row = max((coordinate_to_tuple(key)[0] for key in raw_cells), default=1)
        sheet = book[SHEET]
        headers = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1, max_col=len(HEADERS)))]
        _headers(headers)
        rows = []
        for row_index, record in enumerate(sheet.iter_rows(min_row=2, max_row=max_row, max_col=len(HEADERS)), 2):
            if all(cell.value is None for cell in record):
                continue
            row = {}
            for col_index, (key, cell) in enumerate(zip(headers, record), 1):
                raw = raw_cells.get(f"{get_column_letter(col_index)}{row_index}")
                value = cell.value
                cell_type = cell.data_type
                if raw is not None and raw.attrib.get('t', 'n') == 'n' and raw.find('s:v', NS) is not None and cell_type not in {'f', 'd', 'b', 'e'}:
                    value = raw.find('s:v', NS).text
                    cell_type = 'n'  # Lexical numeric, never Python float.
                elif value is not None:
                    value = str(value)
                if value is not None and len(value) > MAX_TEXT:
                    fail('SPREADSHEET_CELL_TEXT_LIMIT', 413)
                row[key] = {'value': value or '', 'type': cell_type, 'format': cell.number_format}
            rows.append(row)
        return metadata, rows
    except (KeyError, ValueError, StopIteration, TypeError, AttributeError, IndexError, SyntaxError, OSError, InvalidFileException) as exc:
        raise HTTPException(422, 'INVALID_XLSX_STRUCTURE') from exc
    finally:
        if book is not None:
            book.close()


def editable_value(key, cell):
    if cell['type'] not in {'s', 'inlineStr', 'n'}:
        raise ValueError('UNSUPPORTED_CELL_TYPE')
    value = cell['value']
    if value == '':
        return None
    if key not in NUMERIC:
        if cell['type'] == 'n':
            raise ValueError('TEXT_CELL_REQUIRED')
        return value
    percentage = value.endswith('%')
    number = value[:-1] if percentage else value
    if not NUMBER.fullmatch(number) or len(number) > 64:
        raise ValueError('INVALID_DECIMAL_TEXT')
    try:
        with localcontext(Context(prec=50)):
            result = Decimal(number)
            if not result.is_finite():
                raise ValueError('NONFINITE_NUMBER')
            if cell['type'] == 'n' and '%' in cell['format']:
                if key not in PERCENTS or not re.fullmatch(r'0(?:\.0+)?%', cell['format']):
                    raise ValueError('UNSUPPORTED_PERCENT_FORMAT')
                result *= 100
            elif percentage and key not in PERCENTS:
                raise ValueError('PERCENT_NOT_MONEY')
            return str(result)
    except InvalidOperation as exc:
        raise ValueError('INVALID_DECIMAL_TEXT') from exc


class CommercialUploadLimitMiddleware:
    """Bound multipart transport as well as decoded file/archive sizes."""
    def __init__(self, app):
        self.app = app

    async def __call__(self, scope, receive, send):
        if scope['type'] != 'http' or scope['method'] != 'POST' or not scope['path'].endswith('/commercial-exchange/stage'):
            return await self.app(scope, receive, send)
        from starlette.responses import JSONResponse
        from starlette.formparsers import MultiPartException
        limit = MAX_FILE + 65536
        headers = dict(scope.get('headers', []))
        try:
            if int(headers.get(b'content-length', b'0')) > limit:
                return await JSONResponse({'detail': 'SPREADSHEET_FILE_LIMIT'}, status_code=413)(scope, receive, send)
        except ValueError:
            return await JSONResponse({'detail': 'INVALID_CONTENT_LENGTH'}, status_code=400)(scope, receive, send)
        received = 0
        async def limited_receive():
            nonlocal received
            message = await receive()
            received += len(message.get('body', b''))
            if received > limit:
                # MultipartParser closes its spool files when this exception is raised.
                raise MultiPartException('SPREADSHEET_FILE_LIMIT')
            return message
        await self.app(scope, limited_receive, send)
